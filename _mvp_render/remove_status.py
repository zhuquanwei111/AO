from docx import Document
from pathlib import Path
p=Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
d=Document(p)
for para in list(d.paragraphs):
    if para.text.startswith('M0 当前状态（2026-10-11）'):
        para._element.getparent().remove(para._element)
d.save(p)
print('saved',p,'chars=',sum(len(x.text) for x in d.paragraphs)+sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells))
