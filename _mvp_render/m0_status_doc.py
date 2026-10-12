from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from pathlib import Path
p=Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
d=Document(p)
for para in d.paragraphs:
    if para.text.startswith('首期只支持 spec'):
        para.text='首期只支持 spec → plan → build → test → review → ship。A 已确认：DeepSeek 第三方公共 API 用于 build，接受供应商服务端实际路由（包括 V4.1 Flash），记录响应标识；客户端不自动切模。GPT-6.1 Sol 负责规划、架构和只读终审。失败保留现场，不自动修复、不自动全盘回滚。'
    if para.text.startswith('M1 环境基线'):
        para.text='M1 环境基线：M0 通过后安装 WSL2 Ubuntu、OpenCode、Git、Node/Python 和 Playwright MCP；仓库放在 WSL2 内，固定依赖与最小权限配置；在 OpenCode 对 GPT 规划/终审及 DeepSeek build 分别做无副作用连通性核验。未通过则暂停，不进入 M2。'
        anchor=para
# Add current M0 status once, before M1.
if not any(x.text.startswith('M0 当前状态（2026-10-11）') for x in d.paragraphs):
    el=OxmlElement('w:p'); anchor._p.addprevious(el)
    new=Paragraph(el,anchor._parent); new.style=anchor.style
    new.add_run('M0 当前状态（2026-10-11）：本机 WSL2 默认版本为 2，但仅有 docker-desktop 分发版；Ubuntu 与 OpenCode CLI 尚未安装。未带凭据访问天翼与 OpenAI 的 /v1/models 均返回 401，只能证明 HTTPS 端点可达，不能证明模型连通；当前 PowerShell 会话未提供角色独立 API Key。M0 尚未通过，等待角色凭据后再做认证调用、模型标识与限流/超时核验。')
d.save(p)
print('saved',p,'chars=',sum(len(x.text) for x in d.paragraphs)+sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells))
