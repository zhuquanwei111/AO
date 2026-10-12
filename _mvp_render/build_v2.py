from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / '多智能体协同开发系统搭建方案_多Harness优化版_v2.md'
OUTPUT = SOURCE.with_suffix('.docx')
doc = Document(ROOT / '多智能体协同开发系统搭建方案_MVP版_已拍板A.docx')
for element in list(doc._element.body):
    if element.tag != qn('w:sectPr'):
        doc._element.body.remove(element)

sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(17), Mm(17)
sec.left_margin, sec.right_margin = Mm(18), Mm(18)
sec.header_distance, sec.footer_distance = Mm(7), Mm(8)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))):
    sec._sectPr.remove(grid)
for part in [sec.header, sec.footer]:
    for el in list(part._element):
        part._element.remove(el)
    part.add_paragraph()

def set_font(target, size, bold=False, name='Microsoft YaHei'):
    target.font.name = name
    target.font.size = Pt(size)
    target.font.bold = bold
    target.font.italic = False
    target.font.color.rgb = RGBColor(0, 0, 0)
    rp = target._element.get_or_add_rPr()
    fonts = rp.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rp.insert(0, fonts)
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + attr), name)

for name, size, bold in [('Normal',10.5,False),('Title',21,True),('Subtitle',10.5,False),('Heading 1',14,True),('Heading 2',11.5,True),('List Number',10.5,False)]:
    style=doc.styles[name]
    set_font(style,size,bold)
    pf=style.paragraph_format
    pf.space_before=Pt(0)
    pf.space_after=Pt(5)
    pf.line_spacing=Pt({'Title':30,'Subtitle':16,'Heading 1':20,'Heading 2':17}.get(name,14.5))
    pf.widow_control=True
    pf.keep_with_next=name in ['Title','Subtitle','Heading 1','Heading 2']
    pp=style._element.find(qn('w:pPr'))
    if pp is not None:
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
        for tag in ['pBdr','shd']:
            for item in list(pp.findall(qn('w:'+tag))): pp.remove(item)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)

def write_inline(p, text, size=None, bold=False):
    for i, bit in enumerate(re.split(r'(`[^`]+`)',text)):
        code=bit.startswith('`') and bit.endswith('`')
        run=p.add_run(bit[1:-1] if code else bit)
        set_font(run, size or 10.5,bold, 'Consolas' if code else 'Microsoft YaHei')
        if code:
            run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

def add_hyperlink(p,label,url):
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r')
    props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'174A72');props.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');props.append(sz)
    fonts=OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'),'Microsoft YaHei');fonts.set(qn('w:eastAsia'),'Microsoft YaHei');props.append(fonts)
    run.append(props)
    text=OxmlElement('w:t');text.text=label;run.append(text)
    link.append(run);p._p.append(link)

def table(rows, index):
    widths={0:[28,146],1:[31,57,86],2:[29,60,85],3:[31,64,79],4:[37,137]}[index]
    t=doc.add_table(rows=1,cols=len(rows[0]))
    t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,width in zip(t.columns,widths): col.width=Mm(width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        item=OxmlElement('w:'+side)
        item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
    pr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            hdr=OxmlElement('w:tblHeader');trpr.append(hdr)
        for ci,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,value in [('top',75),('bottom',75),('left',95),('right',95)]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),str(value));m.set(qn('w:type'),'dxa');margins.append(m)
            tcpr.append(margins)
            if ri==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E5EDF5');tcpr.append(shd)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.line_spacing=Pt(13.5)
            p.paragraph_format.keep_with_next=ri==0
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 else WD_ALIGN_PARAGRAPH.LEFT
            write_inline(p,text,9.5,ri==0)
    after=doc.add_paragraph()
    after.paragraph_format.space_after=Pt(0)
    after.paragraph_format.space_before=Pt(0)
    after.paragraph_format.line_spacing=1
    after.add_run().font.size=Pt(2)

lines=SOURCE.read_text(encoding='utf-8').splitlines()
i=0;table_index=0;next_page=False
while i<len(lines):
    line=lines[i].strip()
    if not line:
        i+=1;continue
    if line=='<!-- PAGE -->':
        next_page=True;i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',x or '-') for x in row):block.append(row)
            i+=1
        table(block,table_index);table_index+=1;continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):
        p=doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):
        p=doc.add_paragraph(line[4:],style='Heading 2')
    elif i==2:
        p=doc.add_paragraph('多 Harness 协作优化版 2.0    2026 年 10 月 12 日',style='Subtitle')
    elif re.match(r'^\[\d+\]',line) and 'https://' in line:
        label,url=line.split('https://',1)
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=Pt(13.5)
        add_hyperlink(p,label.rstrip('： '),'https://'+url)
    else:
        p=doc.add_paragraph()
        write_inline(p,line)
    if next_page:
        p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0)
        next_page=False
    i+=1

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();set_font(r,9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.append(field)
doc.core_properties.title='多智能体协同开发系统搭建方案'
doc.core_properties.subject='多 Harness 协作优化版 2.0'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.save(OUTPUT)
visible=''.join(doc._element.itertext())
display=''.join(doc._element.xpath('.//w:t/text()'))
original=Document(Path(r'C:\Users\admin\Desktop\和豆包的对话_1011(1).docx'))
original_text='\n'.join([x.text for x in original.paragraphs]+[c.text for t in original.tables for row in t.rows for c in row.cells])
print(json.dumps({'output':OUTPUT.name,'tables':len(doc.tables),'visible_characters':len(display),'original_characters':len(original_text),'ratio':round(len(display)/len(original_text),3)},ensure_ascii=True))
