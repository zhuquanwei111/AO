from docx import Document
from pathlib import Path
path = Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
doc = Document(path)
for table in doc.tables:
    if table.rows and table.rows[0].cells[0].text == '分类' and len(table.rows[0].cells) == 2:
        for row in table.rows[1:]:
            if row.cells[0].text == '复用开源':
                row.cells[1].text = 'OpenCode；WSL2；agent-skills 六阶段模板；Git worktree；Microsoft Playwright MCP。OpenCode 权限机制不替代自研状态机/网关；依赖锁版本并在 WSL2 核验。[1][2][3][4]'
            elif row.cells[0].text == '需自研':
                row.cells[1].text = '外部 SQLite 状态机；MCP 权限网关与审计；按角色隔离的 Worker/凭据启动；Git worktree 生命周期；API 标识/用量采集、限流超时与阻断策略。'
            elif row.cells[0].text == '二期扩展':
                row.cells[1].text = 'DSH 迁移评估；constraints、webperf、code-simplify；DeepSeek/Kimi/GLM 其他模型或自动路由；ZCode 可参考其天翼 Provider 模型目录，但不作 MVP Harness；并行 Agent 与其他 MCP。'
# Correct the explicit source reconciliation and add current price/source links.
for p in doc.paragraphs:
    if p.text.startswith('DeepSeek 路由不是已确认的首期阻断项:') or p.text.startswith('DeepSeek 路由不是已确认的首期阻断项：'):
        p.text = 'DeepSeek 路由不是已确认的首期阻断项：官方 9 月 10 日公告称 deepseek-v4-pro 请求会路由至 V4.1-Flash；但更新日志又称 9 月 14 日后继续提供 V4 Pro API，当前价格页列出 V4-Pro-0813。两组官方资料存在冲突，不能据其推断第三方供应商实际返回；M0 对所选供应商实测请求 ID、response.model、实际服务/计费标识。无法证明底层版本时标记“供应商未验证”，由所选策略决定是否接受，不得静默切换。[7][8]'
    elif p.text.startswith('模型身份与路由是 M0 风险核验项'):
        p.text = '模型身份与路由是 M0 风险核验项。DeepSeek 官方公告与更新日志/当前价格页对 V4 Pro 是否路由、是否继续提供服务的信息不一致；第三方供应商也可能另行映射。[7][8] 请求模型 ID、response.model 与实际底层服务不必相同。若选 A，按约定接受实际返回的 Flash 并记录；若供应商无法提供实际标识，标注“仅供应商声明/未知”，由用户决定是否接受。B/C 同样核验各自映射。任何策略都不得自动换模型。'
    elif p.text.startswith('[7] DeepSeek：'):
        p.text = '[7] DeepSeek API：V4.1-Flash 公告（含 V4 Pro 路由说明） https://api-docs.deepseek.com/zh-cn/news/news260910/'
    elif p.text.startswith('[8] DeepSeek API：'):
        p.text = '[8] DeepSeek API：更新日志与当前模型价格（包含 V4 Pro API 继续提供及 V4-Pro-0813 信息） https://api-docs.deepseek.com/zh-cn/updates/ ; https://api-docs.deepseek.com/zh-cn/quick_start/pricing/'
    elif p.text.startswith('[9] Kimi API：'):
        p.text = '[9] Kimi API：K3 模型定位与 API 价格说明 https://www.kimi.ai/help/kimi-api/api-overview ; https://www.kimi.ai/help/kimi-api/api-pricing'
    elif p.text.startswith('[10] OpenAI API：'):
        p.text = '[10] OpenAI API：GPT-6.1 Sol 模型资料与价格 https://developers.openai.com/api/docs/models/gpt-6.1-sol'
doc.save(path)
print('saved', path)
print('chars=', sum(len(p.text) for p in doc.paragraphs) + sum(len(c.text) for t in doc.tables for row in t.rows for c in row.cells))
