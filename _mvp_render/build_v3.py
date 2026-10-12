from pathlib import Path
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "多智能体协同开发系统搭建方案_Orchestrator试点版_v3.md"
OUTPUT = SOURCE.with_suffix(".docx")

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(17), Mm(17)
section.left_margin, section.right_margin = Mm(18), Mm(18)

def set_font(target, size=10.5, bold=False, color="000000"):
    target.font.name = "Microsoft YaHei"
    target.font.size = Pt(size)
    target.font.bold = bold
    target.font.color.rgb = RGBColor.from_string(color)
    rpr = target._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "Microsoft YaHei")

for name, size, bold in (
    ("Normal", 10.5, False),
    ("Title", 22, True),
    ("Subtitle", 10, False),
    ("Heading 1", 15, True),
    ("Heading 2", 12, True),
    ("List Bullet", 10.5, False),
    ("List Number", 10.5, False),
):
    style = doc.styles[name]
    set_font(style, size, bold)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = Pt(15 if size <= 12 else 20)
    style.paragraph_format.keep_with_next = name in ("Title", "Subtitle", "Heading 1", "Heading 2")
    style.paragraph_format.widow_control = True
doc.styles["Title"].paragraph_format.space_after = Pt(7)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)

def add_link(p, label, url, size):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174A72")
    props.append(color)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(round(size * 2)))
    props.append(font_size)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    props.append(fonts)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    p._p.append(link)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\(https?://[^)]+\))")

def add_inline(p, text, size=10.5, table_cell=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_font(run, size, True)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:].split("](", 1)
            add_link(p, label, url[:-1], size)
        else:
            run = p.add_run(part)
            set_font(run, size)

def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = OxmlElement("w:" + side)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tcpr.append(margins)

def add_table(rows):
    cols = len(rows[0])
    widths = {
        2: [37, 137],
        3: [34, 57, 83],
    }.get(cols, [174 / cols] * cols)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Mm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = Pt(13)
            if ri == 0:
                p.paragraph_format.keep_with_next = True
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "24476B")
                cell._tc.get_or_add_tcPr().append(fill)
                add_inline(p, text, 9.3)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if ri % 2 == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "F3F6F9")
                    cell._tc.get_or_add_tcPr().append(fill)
                add_inline(p, text, 9.1)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = Pt(2)

def parse_table(lines, start):
    block = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r"[:\- ]+", x or "-") for x in row):
            block.append(row)
        i += 1
    add_table(block)
    return i

lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith("|"):
        i = parse_table(lines, i)
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(line[2:], style="Title")
    elif line.startswith("## "):
        p = doc.add_paragraph(line[3:], style="Heading 1")
    elif line.startswith("### "):
        p = doc.add_paragraph(line[4:], style="Heading 2")
    elif i == 2:
        p = doc.add_paragraph(style="Subtitle")
        add_inline(p, line, 10)
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, line[2:])
    elif re.match(r"^\d+\.\s", line):
        p = doc.add_paragraph(style="List Number")
        add_inline(p, re.sub(r"^\d+\.\s", "", line))
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = Pt(15)
        add_inline(p, line)
    i += 1

doc.core_properties.title = "多智能体协同开发系统搭建方案"
doc.core_properties.subject = "Orchestrator 与 WSL2 试点版 3.0"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.save(OUTPUT)
print(f"created={OUTPUT.name}; tables={len(doc.tables)}; paragraphs={len(doc.paragraphs)}")
