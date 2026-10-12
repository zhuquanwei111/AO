from docx import Document
from pathlib import Path
p=Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
d=Document(p)
for para in d.paragraphs:
    if para.text.startswith('首期只支持 spec'):
        para.text='首期只支持 spec → plan → build → test → review → ship。A 已确认：DeepSeek 第三方公共 API 用于 build，接受供应商服务端实际路由（包括 V4.1 Flash），记录响应标识；客户端不自动切模。GPT-6.1 Sol 负责规划、架构和只读终审。失败保留现场，不自动修复、不自动全盘回滚。'
    elif para.text.startswith('M0 前置核验（策略拍板后）'):
        para.text='M0 前置核验（A 已确认）：不依赖 OpenCode 安装，分别验证 GPT 规划、GPT 终审、DeepSeek build 三个角色凭据的 API 连通性；记录请求 ID、response.model、供应商可提供的实际服务标识、token、耗时、状态码和账单。A 接受实测返回 V4.1 Flash；若 response.model 仅回显别名，记录“底层版本未验证”，不因该映射单独阻断。用测试额度验证限流/超时、每阶段最多 2 次重试、401/403 不重试、失败阻断且不切模型。截图显示 ZCode 天翼 Provider 使用 https://ai.ctaigw.cn/v1、OpenAI Chat Completions，目录含 deepseek-v4-pro-0813、deepseek-v4-flash-vision-exp-0817、kimi-k3、glm-5.3、glm-5.3-flash；这只证明 ZCode 可选该目录，不证明 OpenCode 兼容、实际底层版本或账单。M0 通过后进入 M1；M1 再在 WSL2 OpenCode 中分别验证所选模型配置与角色 Key，成功后才进入 M2。另确认仓库及云端模型调用已获授权。'
    elif para.text.startswith('M1 环境基线'):
        para.text='M1 环境基线：M0 通过后安装 WSL2 Ubuntu、OpenCode、Git、Node/Python 和 Playwright MCP；仓库放在 WSL2 内，固定依赖与最小权限配置；在 OpenCode 对 GPT 规划/终审及 DeepSeek build 分别做无副作用连通性核验。未通过则暂停，不进入 M2。'
d.save(p)
print('saved',p,'chars=',sum(len(x.text) for x in d.paragraphs)+sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells))
