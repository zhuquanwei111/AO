from docx import Document
from pathlib import Path
p=Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
d=Document(p)
for para in d.paragraphs:
    if para.text.startswith('推荐 B 作为最稳妥的流程 MVP'):
        para.text='推荐 B 作为最稳妥的流程 MVP：首期固定一个编码模型建立可比较的质量基线，减少 DeepSeek/Kimi 供应商映射和质量波动对首轮验收的影响；但 GPT 编码能力与实际返回版本仍须 M0/M3 验证，不能预设质量不降。代价是成本可能上升并需调整 GPT 编码权限。若优先控制成本且愿意接受 Flash，A 是改动最小的备选；不建议首期同时接入多模型。30% 降本仍是待验证假设。'
    elif para.text.startswith('模型策略：A 接受 DeepSeek 公共 API'):
        para.text='模型策略：A 接受 DeepSeek 公共 API 服务端实际路由并继续，记录请求 ID 和供应商返回标识；B build 使用 GPT-6.1 Sol；C build 使用 Kimi K3。最终仅启用拍板组合，不做客户端自动路由、自动降级或切换。记录 response.model、usage、耗时和状态码；该字段可能只是请求别名回显，标记“已验证/仅回显/未知”，不可据此宣称底层版本已证实。API 连续失败后阶段阻断。'
    elif para.text.startswith('M0 前置核验（策略拍板后）'):
        para.text='M0 前置核验（策略拍板后）：对模型矩阵内每个模型做连通性、请求 ID/响应 model/供应商实际服务标识核验，并记录 token、耗时、状态码；验证供应商限流与超时行为，确认每阶段最多 2 次重试、鉴权失败不重试、失败阻断且不切模型。A 接受实测返回 V4.1 Flash；B 验证 GPT 编码 Worker 独立凭据与网关写权限；C 验证 Kimi K3 映射。截图显示天翼 Provider 目录含 deepseek-v4-pro-0813、deepseek-v4-flash-vision-exp-0817、kimi-k3、glm-5.3、glm-5.3-flash；这只证明 ZCode 可选该目录，不证明 OpenCode 兼容、实际底层版本或账单，须在 OpenCode 单独核验。另确认仓库与云端模型调用已获授权。未通过则不进入 M1。'
    elif para.text.startswith('模型策略拍板后，按角色固定请求模型 ID'):
        para.text='模型策略拍板后，按角色固定请求模型 ID；关闭客户端自动路由、自动降级和跨模型重试。A 可接受供应商服务端路由。OpenCode Agent 只能通过 MCP 权限网关使用文件、终端、Git 和 Playwright；显式关闭绕过网关的直接文件写入与 Shell。[2]'
d.save(p)
print('saved',p,'chars=',sum(len(x.text) for x in d.paragraphs)+sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells))
