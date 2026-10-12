from docx import Document
from pathlib import Path
p=Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
d=Document(p)

def replace_para(prefix, text):
    for para in d.paragraphs:
        if para.text.startswith(prefix):
            if para.runs:
                para.runs[0].text=text
                for run in para.runs[1:]: run.text=''
            else:
                para.add_run(text)
            return True
    return False

changes=[]
def do(prefix,text):
    if not replace_para(prefix,text): raise SystemExit('paragraph not found: '+prefix)
    changes.append(prefix)

do('版本：MVP 1.3', '版本：MVP 1.4（模型策略 A 已确认；GPT 订阅认证已澄清）    日期：2026-10-11    目标：在授权代码仓库验证六阶段交付闭环')
do('每次工具/API 调用记录', '每次模型或工具调用记录任务 ID、Agent/角色、请求模型 ID、供应商实际返回标识、token 用量、耗时、状态码及结果。供应商或订阅接口未暴露的字段标为“不可得”，不得用请求 ID 冒充实际返回模型。审计不记录完整提示、密钥或代码原文，参数先脱敏。第三方 API Key 按角色独立管理，DeepSeek 天翼 Key 仅供 build Worker 使用，不全局共享。GPT 规划/终审通过 OpenCode 官方 OpenAI Provider 的 ChatGPT 订阅登录授权（OpenCode 文档列出 Plus/Pro 登录），不是第三方 API，也不配置 OpenAI API Key；两个只读 GPT 角色可共用经授权的订阅会话，工具只读边界由外部状态机和 MCP 网关强制。订阅授权文件仅限 WSL2 服务用户访问且不入库。超时、限流按 Retry-After/退避规则处理，每阶段最多重试 2 次；鉴权失败（401/403）不重试。仍失败则阶段失败阻断，不自动换模型。网关故障或审计写入失败时拒绝写操作。')
do('数据边界：首期按任务', '数据边界：首期按任务将所需代码上下文发送至指定云端模型；GPT 规划/终审使用 OpenAI 官方 ChatGPT 订阅授权，DeepSeek build 使用天翼第三方 API。首期不实现敏感代码识别、屏蔽或外发审批。MCP 权限网关仍控制工具和写操作；第三方 API Key 按角色隔离，GPT 只读权限由网关执行；调用审计不记录密钥或完整提示内容。')
do('M0 前置核验（A 已确认）', 'M0 前置核验（A 已确认）：GPT 侧确认 ChatGPT 订阅计划和账户可用，并按 OpenCode 官方 OpenAI Provider 流程使用 ChatGPT Plus/Pro 登录授权；此项不是 OpenAI API Key 连通性测试，不调用 api.openai.com API。GPT-6.1 Sol 是否出现在 OpenCode 模型列表、订阅额度及真实请求结果，安排在 M1 安装 OpenCode 后核验。DeepSeek build 侧用独立天翼 API Key 验证认证连通性；请求模型 ID、供应商实际返回标识、token 用量、耗时、状态码和账单逐项记录，A 接受实测返回 V4.1 Flash。若供应商只回显别名或某字段不暴露，明确标注“底层版本未验证”或“不可得”，不伪造信息。使用测试额度验证限流/超时、每阶段最多 2 次重试、401/403 不重试、失败阻断且不切模型。截图显示 ZCode 天翼 Provider 使用 https://ai.ctaigw.cn/v1、OpenAI Chat Completions，目录含 deepseek-v4-pro-0813、deepseek-v4-flash-vision-exp-0817、kimi-k3、glm-5.3、glm-5.3-flash；这只证明 ZCode 可选该目录，不证明 OpenCode 兼容、实际底层版本或账单。M0 完成环境/授权准备后进入 M1；M1 在 WSL2 OpenCode 中验证订阅登录、GPT-6.1 Sol 可用性和 DeepSeek 角色配置，成功后才进入 M2。另确认仓库及云端模型调用已获授权。')
do('M1 环境基线', 'M1 环境基线：完成前置核验后安装 WSL2 Ubuntu、OpenCode、Git、Node/Python 和 Playwright MCP；仓库放在 WSL2 内，固定依赖与最小权限配置。OpenCode 使用官方 OpenAI Provider 的 ChatGPT Plus/Pro 登录授权，核验 GPT-6.1 Sol 是否可选并做无副作用调用；DeepSeek 天翼 API Key 仅配置给 build Worker 并核验连通性。GPT 不配置 OpenAI API Key。未通过则暂停，不进入 M2。')
do('首期不设置敏感代码外发限制', '首期不设置敏感代码外发限制，所需代码上下文按角色发送至云端模型服务：GPT 规划/终审经 OpenAI 官方 ChatGPT 订阅授权，DeepSeek build 经天翼第三方 API。工具权限、DeepSeek build Key 隔离和脱敏调用审计照常执行。')

# Update the worker credential matrix.
t=d.tables[1]
t.rows[4].cells[1].text='按角色固定模型 Worker；凭据按认证方式隔离'
t.rows[4].cells[2].text='GPT：OpenCode 官方 ChatGPT Plus/Pro 订阅登录（非 API Key）；DeepSeek：build 专用天翼 API Key。工具权限由网关按角色控制。'
# Update the reusable / custom boundary row.
t=d.tables[3]
t.rows[2].cells[1].text='外部 SQLite 状态机；MCP 权限网关与审计；角色 Worker 启动；DeepSeek build Key 隔离；GPT 订阅授权接入；Git worktree 生命周期；模型标识/用量采集与限流阻断策略。'
# Update the cost comparison and explicitly separate subscription allocation from API billing.
t=d.tables[4]
t.rows[7].cells[1].text='同一批 30 个匹配任务分别运行 GPT-only 与 MVP。记录通过任务数、质量、GPT 订阅计划与额度消耗、DeepSeek 实际账单；同时报告现金支出和按通过任务分摊的成本。GPT 订阅为固定周期费用，不作为逐 Token API 账单。'
t.rows[8].cells[1].text='MVP/GPT-only ≤ 0.70 仅为待验证假设。分摊口径为测试期订阅费/该组通过任务数 + 该组 DeepSeek 实付/通过任务数。样本不足或质量不等价时不宣称节省。'

# Add source links once, after existing references.
existing='\n'.join(x.text for x in d.paragraphs)
refs=[
    ('[11] OpenCode：OpenAI Provider 的 ChatGPT Plus/Pro 登录 https://opencode.ai/docs/providers/', '[11] OpenCode：OpenAI Provider 的 ChatGPT Plus/Pro 登录 https://opencode.ai/docs/providers/'),
    ('[12] OpenAI：在其他应用中使用 ChatGPT 计划（无需 API Key） https://help.openai.com/en/articles/20001542-using-your-chatgpt-plan-in-other-apps-and-sites', '[12] OpenAI：在其他应用中使用 ChatGPT 计划（无需 API Key） https://help.openai.com/en/articles/20001542-using-your-chatgpt-plan-in-other-apps-and-sites')
]
last=None
for para in d.paragraphs:
    if para.text.startswith('[10]'):
        last=para
if '[11] OpenCode：OpenAI Provider' not in existing:
    if last is None: raise SystemExit('reference [10] not found')
    p11=last.insert_paragraph_before(refs[0][1])
    p12=last.insert_paragraph_before(refs[1][1])
    # Move the new paragraphs after [10] by reordering their XML nodes.
    body=last._p.getparent()
    body.remove(p11._p); body.remove(p12._p)
    body.insert(body.index(last._p)+1,p11._p)
    body.insert(body.index(p11._p)+1,p12._p)
    changes.append('reference [11]-[12]')

# Correct stale wording in the KPI explanation, if present.
for para in d.paragraphs:
    if para.text.startswith('成本比率 ='):
        para.runs[0].text='成本比率按同一测试期、同一任务质量门槛比较；GPT 订阅是固定费用，DeepSeek 按供应商账单计费。30% 仅为待验证假设，不作为已达成承诺。'
        for run in para.runs[1:]: run.text=''
        changes.append('cost ratio paragraph')

d.save(p)
print('saved',p)
print('changes:',', '.join(changes))
