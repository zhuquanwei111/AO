from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from pathlib import Path

path = Path(r'C:\Users\admin\Documents\ChatGPT\多智能体搭建\多智能体协同开发系统搭建方案_MVP版.docx')
doc = Document(path)

replacements = {
'版本：MVP 1.1    日期：2026-10-11    目标：在授权代码仓库验证六阶段交付闭环': '版本：MVP 1.2（模型策略待拍板）    日期：2026-10-11    目标：在授权代码仓库验证六阶段交付闭环',
'首期只支持 spec → plan → build → test → review → ship。GPT-6.1 Sol 负责规划、架构和只读终审；DeepSeek V4 Pro 仅负责编码。测试由确定性命令与 Playwright 执行。失败时保留任务现场，不自动修复、不自动全盘回滚。': '首期只支持 spec → plan → build → test → review → ship。build 编码模型在 A/B/C 中拍板；当前推荐 B（GPT-6.1 Sol），待用户确认。规划/架构/终审 GPT Worker 保持只读，若选 B，另设独立凭据和受网关约束的 GPT 编码 Worker。失败保留现场，不自动修复、不自动全盘回滚。',
'重要前置条件：DeepSeek 官方公告当前说明，公共 API 对 deepseek-v4-pro 的请求在 V4.1-Pro 上线前会路由到 V4.1-Flash。若必须使用真实 V4 Pro，启动前必须取得能保证模型版本的独立端点；无法保证时阻止 build 阶段，不静默替换模型。[7]': 'DeepSeek 路由不是已确认的首期阻断项：历史公告与当前模型定价/目录信息存在版本差异，第三方供应商也可能有自己的映射。M0 必须实测请求 ID、响应 model 字段及供应商可提供的实际服务/计费标识；不能证明底层版本时明确标为“供应商未验证”，由所选策略决定是否接受，不得静默切换。[7][8]',
'首期模型请求均固定角色和模型 ID；关闭自动模型选择、模型降级及跨模型重试。OpenCode 默认工具权限较宽，部署时必须显式拒绝直接文件修改和 Shell 工具，并只开放网关 MCP 工具。[2]': '模型策略拍板后，按角色固定请求模型 ID；关闭自动路由、自动降级和跨模型重试。OpenCode Agent 只能通过 MCP 权限网关使用文件、终端、Git 和 Playwright；显式关闭绕过网关的直接文件写入与 Shell。[2]',
'每次工具调用记录任务 ID、Agent/角色、工具名、资源路径或命令摘要、时间、结果码和关联模型/用量。日志不保存源码正文、密钥或完整敏感参数；调用结果中的密钥模式必须脱敏。网关故障或审计写入失败时拒绝执行写操作。': '每次工具/API 调用记录任务 ID、Agent/角色、工具名或请求模型 ID、供应商返回 model 标识、token 用量、耗时、状态码及结果。审计不记录完整提示、密钥或代码原文；参数先脱敏。API Key 按角色独立管理，不全局共享。超时、限流按供应商 Retry-After/退避规则处理，每阶段最多重试 2 次；鉴权失败（401/403）不重试。仍失败则阶段失败阻断，不自动换模型。网关故障或审计写入失败时拒绝写操作。',
'模型固定：OpenAI 使用 gpt-6.1-sol；DeepSeek 使用经端点验证的 V4 Pro 版本。记录请求模型 ID 与供应商返回的版本信息；无版本证据、限流、超时或鉴权失败时将任务置为失败阻断，不自动切换模型。': '模型策略：A 接受 DeepSeek 公共 API 实际路由并继续，记录请求 ID 和供应商返回标识；B build 使用 GPT-6.1 Sol；C build 使用 Kimi K3。最终仅启用拍板组合，不自动路由或切换。记录 response.model、usage、耗时和状态码；该字段可能只是请求别名回显，必须标记“已验证/仅回显/未知”，不可据此宣称底层版本已证实。API 连续失败后阶段阻断。',
'M0 前置核验：确认 GPT-6.1 Sol 的 OpenCode 接入方式；取得可证明实际 V4 Pro 的编码端点；确认试点仓库及模型调用已获授权。三项任一未通过，暂停模型编码接入。': 'M0 前置核验（策略拍板后）：对模型矩阵内每个模型做连通性、请求 ID/响应 model/供应商实际服务标识核验，并记录 token、耗时、状态码；验证供应商限流与超时行为，确认每阶段最多 2 次重试、鉴权失败不重试、失败阻断且不切模型。A 接受实测返回 V4.1 Flash；B 验证 GPT 编码 Worker 独立凭据与网关写权限；C 验证 Kimi K3 映射。截图可确认 ZCode 的天翼第三方 Provider 及其模型目录，不证明 OpenCode 兼容、实际底层版本或账单；须在 OpenCode 单独核验。另确认仓库与云端模型调用已获授权。未通过则不进入 M1。',
'M2 控制与权限：实现 SQLite 状态机、阶段产物结构、审批入口、角色 Worker、MCP 网关、审计和模型版本核验；先验证越权调用会被拒绝。': 'M2 控制与权限：实现 SQLite 状态机、阶段产物、审批入口、角色 Worker、MCP 网关、审计及模型标识记录/校验；先验证越权调用会被拒绝。',
'DeepSeek V4 Pro 版本锁定是首期上线阻断项。DeepSeek 官方公告当前说明公共 API 的 deepseek-v4-pro 会路由至 V4.1-Flash，直至 V4.1-Pro 上线。[7] 首期不得把模型 ID 等同实际模型版本。若供应商不能提供固定版本端点及可审计的响应标识，状态机禁止进入 build；由项目负责人决定等待、采购独立托管端点或调整范围。': '模型身份与路由是 M0 风险核验项，不预设 DeepSeek 公共 API 一定返回某个版本。官方资料有时间差异，且第三方供应商可自行映射；请求模型 ID、response.model 与实际底层服务不必相同。[7][8] 若选 A，按约定接受实际返回的 Flash 并记录；若供应商无法提供实际标识，标注“仅供应商声明/未知”，由用户决定是否接受。B/C 同样需核验各自供应商映射。任何策略都不得自动换模型。',
'首期不设置敏感代码外发限制，按阶段和角色权限调用 GPT-6.1 Sol 与 DeepSeek V4 Pro。工具权限、独立凭据和调用审计照常执行。': '首期不设置敏感代码外发限制，所需代码上下文按角色发送至拍板后的第三方云端模型；工具权限、角色独立凭据和脱敏调用审计照常执行。',
'[7] DeepSeek：V4.1-Flash 公告及 V4 Pro API 路由说明  官方资料': '[7] DeepSeek：V4.1-Flash 公告及 API 变更记录（历史路由说明） https://api-docs.deepseek.com/news/news260910/',
}
for p in doc.paragraphs:
    if p.text in replacements:
        p.text = replacements[p.text]

# Insert concise decision comparison before the architecture section.
anchor = next(p for p in doc.paragraphs if p.text == '3. 运行结构与权限边界')
new_texts = [
    '模型策略拍板项（推荐 B，待用户确认）',
    'A｜接受 DeepSeek 公共 API 路由：编码质量取决于实际返回的 V4 Pro/Flash，不能仅凭请求名推断；成本有望较低，但以天翼等供应商账单为准；接入改动最少、最利于进度；风险是映射和质量波动。接受 Flash 即继续，记录实际返回标识，不因版本名不符而自动切换。',
    'B｜首期编码全部用 GPT-6.1 Sol：模型身份与质量基线较易统一，适合先验证流程；成本预期高于 DeepSeek Flash，须按供应商实账比较；落地难度低至中，需新增独立 GPT 编码 Worker/凭据并经网关限权；供应商可用性与费用仍须实测。此选项明确调整“GPT 只读、不编码”的原规则：规划/终审仍只读，编码使用隔离 Worker。',
    'C｜改用 Kimi K3 编码：官方将 K3 定位用于长程编码/软件工程，但本仓库质量须实测；成本按第三方报价核算（官方直连价仅供参考）；需新增模型映射、凭据和验收，难度中等；实际路由与质量尚无本项目数据，可能增加 M0/M3 时间。[9]',
    '推荐 B 作为最稳妥的流程 MVP：先固定成熟主力模型建立质量基线，避免 DeepSeek/Kimi 供应商映射不透明干扰首轮验收；代价是成本可能上升并需接受 GPT 编码 Worker 的角色规则调整。若优先控制成本且愿意接受 Flash，A 是改动最小的备选；不建议首期同时接入多模型。无论选项，30% 降本均是待验证假设。',
]
for i, txt in enumerate(new_texts):
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    p = Paragraph(element, anchor._parent)
    if i == 0:
        p.style = doc.styles['Heading 2']
    else:
        p.style = doc.styles['Normal']
    p.add_run(txt)

# Update table cells by exact old value.
cell_replacements = {
    'GPT-6.1 Sol 规划/架构/只读评审；DeepSeek V4 Pro 编码': '规划/终审 GPT-6.1 Sol 只读；build 选 A DeepSeek、B GPT-6.1 Sol、C Kimi K3（待拍板）',
    'V4.1 自动路由、Kimi Code、ZCode': 'V4.1 自动模型路由、Kimi Code、ZCode Harness；GLM 与并行 Agent',
    '按角色独立进程与凭据 || GPT 只读代码；DeepSeek 仅在已批准任务 worktree 编码；不共享环境变量凭据': '按角色独立 Worker 与凭据 || 规划/终审 GPT 只读；build 按 A/B/C 固定模型；B 时 GPT 编码 Worker 与只读 Worker 分离；不共享 Key',
    'build || DeepSeek V4 Pro || 按批准计划逐切片改动，仅写任务 worktree；提交 diff 和实际模型标识。模型/端点不匹配则阻断。': 'build || 用户拍板的 A/B/C 编码模型 || 按批准计划逐切片改动，仅写任务 worktree；记录请求 ID、response.model、token、耗时、状态码。A 接受实际路由；不自动切模。',
    '|| DeepSeek V4 Pro 仅在已批准任务 worktree 编码；不共享环境变量凭据': '|| build 按 A/B/C 固定模型；GPT 只读规划/评审与 GPT 编码 Worker 分开；API Key 按角色隔离',
    '基线：同一 30 个匹配任务全部使用 GPT-6.1 Sol；MVP 组：GPT 规划/评审 + DeepSeek 编码。比较每个验收通过任务的输入、输出、推理及缓存 token 和实际账单；报告质量差异。': '基线：同一 30 个匹配任务全部使用 GPT-6.1 Sol；MVP 组使用拍板后的 A/B/C 组合。逐任务记录模型请求 ID、返回标识、token、耗时及第三方实际账单，比较每个验收通过任务成本并报告质量差异。',
    '复用开源 || OpenCode；WSL2；agent-skills 中六阶段对应 Skill/参考文档；Git；Microsoft Playwright MCP；OpenCode MCP/Agent 权限与插件机制。具体依赖需锁版本并做兼容核验。[1][2][3][4]': '复用开源 || OpenCode；WSL2；六阶段 Skill 模板；Git worktree；Playwright MCP。OpenCode 权限机制不替代外部状态机/网关，所有依赖须锁版本并在 WSL2 核验。[1][2][3][4]',
    '二期扩展 || DSH 迁移评估；constraints、webperf、code-simplify；Kimi Code、ZCode；V4.1 或其他模型路由；并行 Agent；数据库/Figma/Web MCP；远程任务和管理界面。': '二期扩展 || DSH 迁移评估；constraints、webperf、code-simplify；按试点评估 DeepSeek/Kimi/GLM 新模型或自动路由；ZCode 可参考其天翼 Provider 目录但不作 MVP Harness；并行 Agent 与其他 MCP。',
}
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text in cell_replacements:
                cell.text = cell_replacements[cell.text]

# Exact table cell edits, preserving the rest of the table wording.
for table in doc.tables:
    for row in table.rows:
        if len(row.cells) == 3 and row.cells[0].text == '阶段' and row.cells[1].text == '执行者' and row.cells[2].text == '产物与退出条件':
            row.cells[3-2].text = '执行者' # no-op, defensive
        if len(row.cells) == 3 and row.cells[0].text == 'build':
            row.cells[1].text = '拍板模型 A/B/C'
            row.cells[2].text = '按批准计划逐切片改动，仅写任务 worktree；记录请求 ID、response.model、token、耗时、状态码。A 接受实际路由；失败阻断，不自动切模。'
        if len(row.cells) == 2 and row.cells[0].text == 'Token 成本对比':
            row.cells[1].text = '基线：同一 30 个匹配任务全部使用 GPT-6.1 Sol；MVP 组使用拍板后的 A/B/C 组合。逐任务记录请求模型 ID、返回标识、token、耗时及第三方实际账单，比较每个验收通过任务成本并报告质量差异。'
        if len(row.cells) == 3 and row.cells[0].text == '模型':
            row.cells[1].text = '规划/终审 GPT-6.1 Sol 只读；build 选 A DeepSeek、B GPT-6.1 Sol、C Kimi K3（待拍板）'
            row.cells[2].text = 'V4.1 自动模型路由、Kimi Code、ZCode Harness；GLM 与并行 Agent'
        if len(row.cells) == 3 and row.cells[0].text == '模型 Worker':
            row.cells[1].text = '按角色独立 Worker 与凭据'
            row.cells[2].text = '规划/终审 GPT 只读；build 按 A/B/C 固定模型；B 时 GPT 编码 Worker 与只读 Worker 分离；不共享 Key'

# Add concise references for comparison assertions.
last_ref = next(p for p in doc.paragraphs if p.text.startswith('[7] DeepSeek:') or p.text.startswith('[7] DeepSeek：'))
extra_refs = [
    '[8] DeepSeek API：当前模型与价格表（与历史路由公告按日期核对） https://api-docs.deepseek.com/zh-cn/quick_start/pricing/',
    '[9] Kimi API：模型定位与平台价格 https://www.kimi.ai/help/kimi-api/api-overview ; https://platform.kimi.ai/',
    '[10] OpenAI API：GPT-6.1 Sol 模型资料 https://developers.openai.com/api/docs/models/gpt-6.1-sol',
]
for txt in extra_refs:
    element = OxmlElement('w:p')
    last_ref._p.addnext(element)
    p = Paragraph(element, last_ref._parent)
    p.style = doc.styles['Normal']
    p.add_run(txt)
    last_ref = p

doc.save(path)
print(f'updated={path}')
print(f'paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}')
print('chars=', sum(len(p.text) for p in doc.paragraphs) + sum(len(c.text) for t in doc.tables for row in t.rows for c in row.cells))
